import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_qa_workbook():
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # STYLES DEFINITIONS
    # -------------------------------------------------------------
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Calibri", size=10, color="000000")
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")
    
    fill_header_s1 = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Navy Blue
    fill_header_s2 = PatternFill(start_color="203764", end_color="203764", fill_type="solid") # Deep Slate Blue
    
    fill_pass = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid") # Light Green
    fill_fail = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid") # Light Orange
    
    fill_high = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid") # Soft Red
    fill_med = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Soft Yellow
    fill_low = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Soft Green
    
    font_high = Font(name="Calibri", size=10, bold=True, color="900000")
    font_med = Font(name="Calibri", size=10, bold=True, color="8A6D00")
    font_low = Font(name="Calibri", size=10, bold=True, color="274E13")
    font_pass = Font(name="Calibri", size=10, bold=True, color="1E4620")
    
    border_thin = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )
    
    align_center = Alignment(horizontal="center", vertical="top", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # -------------------------------------------------------------
    # SHEET 1: Product-Level Testing
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Product-Level Testing"
    ws1.views.sheetView[0].showGridLines = True
    
    headers_s1 = [
        "ID",
        "Module / Page",
        "Feature / Flow",
        "Issue Type",
        "Scenario / Context",
        "Preconditions",
        "Steps to Reproduce",
        "Expected Result",
        "Actual Result",
        "User / Business Impact",
        "Severity",
        "Evidence",
        "Suggested Improvement / Fix",
        "Tester Reasoning"
    ]
    
    ws1.append(headers_s1)
    ws1.row_dimensions[1].height = 28
    
    for col_idx, h in enumerate(headers_s1, 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header_s1
        cell.alignment = align_header
        cell.border = border_thin

    data_s1 = [
        [
            "PL-01",
            "Recruitment",
            "Add Candidate Registration",
            "Functional Bug",
            "Contact Number field accepts alphabetic and arbitrary special characters without validation and successfully persists corrupted candidate data.",
            "Logged in as Admin; navigated to Recruitment > Add Candidate.",
            "1. Enter valid First Name and Last Name.\n2. Enter valid Email.\n3. In 'Contact Number' field, enter non-numeric string: 'INVALID_PHONE_NUMBER_TEXT'.\n4. Click 'Save' button.",
            "System should validate that contact numbers contain only valid digits/plus signs, or display inline validation error 'Allows numbers only' and prevent saving.",
            "System displays 'Success / Successfully Saved' toast notification and persists corrupted alphabetic data into the candidate database record.",
            "Recruiters relying on candidate profiles receive corrupted contact information. Automated SMS/call integrations and HR follow-ups will fail silently.",
            "Medium",
            "evidence/recruitment_phone_saved_behavior.png\nevidence/recruitment_phone_letters.png",
            "Add regex pattern validation on contact number input (e.g., ^[+]?[0-9\\s\\-()]{7,20}$) with clear inline error message 'Please enter a valid phone number'.",
            "Contact numbers are primary communication channels for recruiters. Allowing free-form non-numeric data degrades database integrity and breaks communication workflows."
        ],
        [
            "PL-02",
            "PIM",
            "Employee Search & Filter Reset",
            "UX Issue",
            "Clicking 'Reset' clears search filter input fields but does not reload or refresh the employee table to show the complete dataset.",
            "Logged in as Admin; navigated to PIM > Employee List.",
            "1. Enter non-existent Employee ID or Name filter and click 'Search' (table displays zero or filtered subset).\n2. Click the 'Reset' button.\n3. Observe the search input fields and the employee table below.",
            "Clicking 'Reset' should clear all filter inputs AND automatically refresh the table data to show the default full employee listing (162 records).",
            "Filter inputs are cleared, but the table remains in the filtered/empty state until the user manually clicks 'Search' again with empty fields.",
            "HR users can easily be misled into believing the table currently reflects all employees when in fact it is still displaying an outdated filtered subset.",
            "Medium",
            "evidence/pim_reset_behavior.png\nevidence/pim_after_reset.png",
            "Bind the Reset button action to both clear Vue reactive state models AND trigger a table data re-fetch with default pagination parameters.",
            "In enterprise HR portals, 'Reset' universally implies returning the entire view to its pristine default state. Leaving the table decoupled from the reset action creates cognitive dissonance."
        ],
        [
            "PL-03",
            "Authentication / Session",
            "Post-Logout Browser Back Navigation",
            "Logical Issue",
            "Navigating backward via browser 'Back' button after logging out exposes cached dashboard metrics and structure without re-authenticating.",
            "Active Admin session logged in on OrangeHRM.",
            "1. Log into OrangeHRM dashboard.\n2. Click User Dropdown > Logout.\n3. Verify landing on /auth/login.\n4. Click browser 'Back' button.",
            "Application should intercept back navigation via Cache-Control headers (no-cache, no-store, must-revalidate) and immediately redirect to /auth/login or render an unauthorized session error.",
            "Browser renders cached Dashboard view displaying employee presence statistics, pending actions, and quick-launch widgets before navigating.",
            "On shared or kiosk office computers, subsequent unauthorized users can view confidential company headcount and operational metrics by simply clicking 'Back'.",
            "Medium",
            "evidence/auth_back_navigation.png\nevidence/auth_logout.png",
            "Implement strict server response headers (Cache-Control: no-store, max-age=0, must-revalidate; Pragma: no-cache) and client-side router navigation guards checking token validity on every popstate event.",
            "HR portals handle strictly confidential personnel data. Session invalidation must prevent both server communication and client-side history snooping."
        ],
        [
            "PL-04",
            "PIM",
            "Employee Search with Alphanumeric Query",
            "UX Issue",
            "Searching Employee ID with alphanumeric or special characters triggers inconsistent floating error toast alongside table zero-state.",
            "Logged in as Admin; navigated to PIM > Employee List.",
            "1. In Employee Id filter, enter alphanumeric text: 'TEST_ABC_123'.\n2. Click 'Search' button.\n3. Observe toast notifications and table body.",
            "System should gracefully display standard zero-record table state '(0) Records Found' with 'No Records Found' empty container without firing system error toasts.",
            "A temporary floating toast notification 'No Records Found' appears in the top right while the table also displays 'No Records Found', causing UI redundancy and alert fatigue.",
            "Redundant warning banners create friction and make users question whether the query failed due to a system crash or simply had zero matches.",
            "Low",
            "evidence/pim_search_alphanumeric_error.png\nevidence/pim_search_numeric_no_records.png",
            "Consolidate zero-result feedback exclusively to the table header count span and table container body rather than triggering toast popups.",
            "Toasts are intended for asynchronous system notifications, warnings, or errors. Query results should be confined to the main content container."
        ],
        [
            "PL-05",
            "PIM",
            "Add Employee Name Length Boundary",
            "Improvement",
            "First Name and Last Name inputs enforce hard 30-character limit via HTML maxlength without visual counter or user feedback.",
            "Logged in as Admin; navigated to PIM > Add Employee.",
            "1. In First Name field, attempt to type or paste a name longer than 30 characters (e.g. 50 characters).\n2. Observe input behavior.",
            "Input should either display character counter (e.g., 28/30) or show an informative tooltip/helper text explaining the 30-character maximum.",
            "Input silently truncates text at 30 characters without providing any indication to the user that characters were omitted.",
            "Employees with compound, hyphenated, or long cultural names may have their names unintentionally truncated without HR noticing before saving.",
            "Low",
            "evidence/pim_name_length_limit.png",
            "Add visual character counter badge and clear helper hint 'Max 30 characters' below the name inputs.",
            "Silent input truncation is a known usability anti-pattern that leads to corrupted personal records."
        ],
        [
            "PL-06",
            "Leave Management",
            "Apply Leave Date Validation",
            "Observation",
            "Date range inversion logic correctly flags 'To date should be after from date' and prevents invalid leave applications.",
            "Logged in as Admin; navigated to Leave > Apply.",
            "1. Select Leave Type.\n2. Set 'From Date' = 2026-10-15.\n3. Set 'To Date' = 2026-10-10 (earlier date).\n4. Click form Header or Apply button.",
            "System should flag logical date inversion and disable/prevent submission.",
            "System displays clear red inline validation 'To date should be after from date' under the To Date picker and blocks form submission.",
            "Positive validation: Prevents negative leave duration calculations and database corruption in employee leave balances.",
            "Low",
            "evidence/leave_date_validation_verified.png\nevidence/leave_submit_inverted_dates.png",
            "Maintain existing robust date comparison logic. Consider auto-updating To Date to match From Date when From Date is chosen.",
            "Verified positive implementation of date range business rules."
        ],
        [
            "PL-07",
            "Admin / User Management",
            "System User Creation Validation",
            "Observation",
            "Duplicate username detection ('Admin') and minimum username boundary (<5 characters) properly enforced with instant feedback.",
            "Logged in as Admin; navigated to Admin > User Management > Users > Add.",
            "1. Enter 'adm' (< 5 chars) in Username field -> observe inline message.\n2. Enter 'Admin' (existing username) in Username field -> observe inline message.",
            "System should enforce >=5 characters and check uniqueness against database.",
            "System displays 'Should be at least 5 characters' for short inputs and 'Already exists' for duplicate usernames.",
            "Positive validation: Prevents user account collisions and weak login identifiers across the enterprise.",
            "Low",
            "evidence/admin_username_boundary.png\nevidence/admin_duplicate_username.png",
            "Maintain current responsive debounced API validation check.",
            "Verified positive implementation of security and integrity boundaries for administrative credentials."
        ],
        [
            "PL-08",
            "Recruitment",
            "Candidate Email Format Validation",
            "Observation",
            "Malformed email addresses are strictly validated with descriptive inline guidance.",
            "Logged in as Admin; navigated to Recruitment > Add Candidate.",
            "1. Enter 'invalid-candidate-email-at-domain' in Email field.\n2. Click out of the field (blur).",
            "System should validate standard email RFC pattern and show clear error message.",
            "System displays red inline validation 'Expected format: admin@example.com' and highlights input.",
            "Positive validation: Prevents invalid contact records from entering recruitment talent pipeline.",
            "Low",
            "evidence/recruitment_email_format_error.png",
            "Maintain current format validator.",
            "Verified positive implementation of regex validation on mission-critical communication fields."
        ],
        [
            "PL-09",
            "PIM",
            "Bulk Employee Selection & Action Bar",
            "Observation",
            "Header checkbox accurately selects all rendered employee cards and contextualizes the 'Delete Selected' action bar.",
            "Logged in as Admin; navigated to PIM > Employee List.",
            "1. Click the master checkbox in the table header.\n2. Observe row checkboxes and action buttons above the table.",
            "Header checkbox should toggle all 50 active row checkboxes and reveal 'Delete Selected' batch action button.",
            "All 50 row checkboxes become active and orange 'Delete Selected' button appears with accurate selection count.",
            "Positive validation: Streamlines bulk HR administrative workflows like department reorganizations or record archiving.",
            "Low",
            "evidence/pim_bulk_selection.png",
            "Maintain existing bulk selection mechanism; consider adding 'Select all 162 records across all pages' option for large datasets.",
            "Verified positive implementation of bulk data management UX."
        ],
        [
            "PL-10",
            "Authentication",
            "Mandatory Credential Enforcement",
            "Observation",
            "Empty login submissions trigger simultaneous field validation highlighting required fields.",
            "Logged out on OrangeHRM login page.",
            "1. Leave Username and Password blank.\n2. Click 'Login' submit button.",
            "System should highlight both fields with 'Required' inline message without reloading page.",
            "Both inputs display red borders and 'Required' helper text below inputs simultaneously.",
            "Positive validation: Saves server roundtrips and provides immediate feedback to the user.",
            "Low",
            "evidence/auth_empty_credentials.png",
            "Maintain current client-side required field enforcement.",
            "Verified positive baseline authentication validation."
        ]
    ]

    for row_idx, r_data in enumerate(data_s1, 2):
        ws1.row_dimensions[row_idx].height = 65
        for col_idx, val in enumerate(r_data, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_body
            cell.border = border_thin
            if col_idx in [1, 4, 11]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            
            # Severity coloring
            if col_idx == 11:
                if val == "High":
                    cell.fill = fill_high
                    cell.font = font_high
                elif val == "Medium":
                    cell.fill = fill_med
                    cell.font = font_med
                elif val == "Low":
                    cell.fill = fill_low
                    cell.font = font_low

    # -------------------------------------------------------------
    # SHEET 2: Automation Testing
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Automation Testing")
    ws2.views.sheetView[0].showGridLines = True
    
    headers_s2 = [
        "Test Case ID",
        "Test Case Name",
        "User Flow / Feature",
        "Tool",
        "Framework",
        "Test Type",
        "Preconditions",
        "Expected Result",
        "Actual Result",
        "Execution Status",
        "Bugs / Inconsistencies",
        "Regression Risk",
        "Risk of Missing Coverage",
        "Further Automation Recommendation",
        "Notes / Evidence"
    ]
    
    ws2.append(headers_s2)
    ws2.row_dimensions[1].height = 28
    
    for col_idx, h in enumerate(headers_s2, 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header_s2
        cell.alignment = align_header
        cell.border = border_thin

    data_s2 = [
        [
            "TC01",
            "Valid User Login Flow",
            "Authentication / Session",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Smoke",
            "SauceDemo application accessible at https://www.saucedemo.com/",
            "User is authenticated, redirected to /inventory.html, and 'Products' header is displayed.",
            "Redirected to inventory.html, title displayed as 'Products'.",
            "PASS",
            "None. Clean login transition.",
            "High. If login fails, all downstream e-commerce flows are completely blocked.",
            "Critical. Unauthenticated users cannot view products or make purchases.",
            "Automate performance response time under load and multi-browser cross-platform matrix.",
            "evidence/saucedemo/tc01_valid_login_success.png"
        ],
        [
            "TC02",
            "Invalid Credentials Validation",
            "Authentication / Security",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "Application on login screen.",
            "Error banner displays 'Epic sadface: Username and password do not match any user in this service' and user remains on login page.",
            "Expected error banner displayed with exact error text; stayed on login URL.",
            "PASS",
            "None. Validation correctly enforced.",
            "Medium. Bad credentials could accidentally grant access or reveal sensitive stack traces.",
            "Moderate. Prevents account enumeration and unauthorized access.",
            "Automate SQL injection strings, XSS payload inputs, and password field masking verification.",
            "evidence/saucedemo/tc02_invalid_credentials_error.png"
        ],
        [
            "TC02b",
            "Locked Out User Error Flow",
            "Authentication / Account Security",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "Application on login screen.",
            "Error banner displays 'Epic sadface: Sorry, this user has been locked out.'",
            "Error banner correctly shown for locked_out_user.",
            "PASS",
            "None. Account state status accurately handled.",
            "Medium. Inactive or locked accounts must not access active storefront inventory.",
            "Moderate. Security regression if locked accounts bypass authorization filters.",
            "Automate account unlock triggers and admin override simulations.",
            "evidence/saucedemo/tc02_locked_out_user.png"
        ],
        [
            "TC02c",
            "Empty Username Field Validation",
            "Authentication / Form Validation",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "Application on login screen.",
            "Inline error displays 'Epic sadface: Username is required'.",
            "Error banner displayed accurately with required prompt.",
            "PASS",
            "None. Client validation working as designed.",
            "Low. Minor user guidance regression.",
            "Low. Basic field requirement validation.",
            "Automate empty password field and whitespace-only username variations.",
            "evidence/saucedemo/tc02_empty_username.png"
        ],
        [
            "TC03",
            "Product Listing Inventory Display",
            "Catalog / Inventory Display",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "Logged in as standard_user on inventory page.",
            "Inventory grid renders exactly 6 product cards with non-empty titles, valid descriptions, and prices.",
            "All 6 items rendered with complete metadata.",
            "PASS",
            "None. Catalog layout structurally sound.",
            "High. Missing catalog items directly cause lost sales revenue.",
            "High. Product discovery is the core e-commerce funnel stage.",
            "Automate dynamic product inventory count validation against backend API catalog.",
            "evidence/saucedemo/tc03_product_listing.png"
        ],
        [
            "TC03b",
            "Product Item Card Structure & CTA",
            "Catalog / Product Components",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "Logged in as standard_user on inventory page.",
            "Each product card contains a title, price formatted with currency symbol '$', and an active 'Add to cart' button.",
            "All 6 cards verified with valid title, '$' price prefix, and 'Add to cart' CTA.",
            "PASS",
            "None. Card elements properly populated.",
            "High. Missing prices or broken add-to-cart buttons halt the buying flow.",
            "High. Ensures consistent pricing presentation across catalog.",
            "Automate currency localization checks (e.g. EUR, GBP) and image rendering integrity.",
            "evidence/saucedemo/tc03_product_listing.png"
        ],
        [
            "TC03c",
            "Product Sorting: Price Low to High",
            "Catalog / Sort Functionality",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Regression",
            "Logged in as standard_user on inventory page.",
            "Selecting 'Price (low to high)' rearranges items in ascending price order [7.99, 9.99, 15.99, 15.99, 29.99, 49.99].",
            "Price array matched sorted array exactly [7.99, 9.99, 15.99, 15.99, 29.99, 49.99].",
            "PASS",
            "None. Sort algorithm operates correctly.",
            "Medium. Broken sorting frustrates shoppers looking for budget options.",
            "Medium. Ensures sort algorithms handle ties and floating point values properly.",
            "Automate remaining sort options: Name (A to Z), Name (Z to A), and Price (high to low).",
            "evidence/saucedemo/tc03c_sort_price_low_high.png"
        ],
        [
            "TC04",
            "Add Single Product to Cart & Badge Sync",
            "Cart / State Management",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "Logged in on inventory page with empty cart.",
            "Cart badge increments to '1', and clicked product button text transforms from 'Add to cart' to 'Remove'.",
            "Badge showed '1'; button text updated to 'Remove'.",
            "PASS",
            "None. Reactive cart state synced successfully.",
            "High. If cart badge fails to update, user is unsure if item was added and may double-purchase.",
            "High. Core state management regression.",
            "Automate rapid double-click handling and add-to-cart latency.",
            "evidence/saucedemo/tc04_add_single_item.png"
        ],
        [
            "TC04b",
            "Add Multiple Products to Cart",
            "Cart / Multi-item State",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "Logged in on inventory page with empty cart.",
            "Adding 2 distinct items updates cart badge count to '2'.",
            "Badge successfully updated to '2' upon adding second item.",
            "PASS",
            "None. Multi-item state accumulation verified.",
            "High. Failure in multi-item cart accumulation restricts cart basket size.",
            "Medium. Prevents quantity desynchronization across catalog.",
            "Automate adding all 6 catalog items and boundary cart limits.",
            "evidence/saucedemo/tc04_add_two_items.png"
        ],
        [
            "TC05",
            "Cart Content & Item Integrity Validation",
            "Cart / Item Review",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "2 items added to cart; navigated to /cart.html.",
            "Cart page renders exactly 2 items with matching product names ('Sauce Labs Backpack', 'Sauce Labs Bike Light') and prices.",
            "Cart displayed 2 items matching names added from inventory.",
            "PASS",
            "None. Cart contents matched selected products.",
            "High. Showing incorrect items or prices in cart destroys customer trust and causes checkout abandonment.",
            "High. Cart accuracy is critical before entering payment funnel.",
            "Automate price sum verification and persistence across page reloads.",
            "evidence/saucedemo/tc05_cart_content.png"
        ],
        [
            "TC05b",
            "Remove Item from Cart Flow",
            "Cart / Modification",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "1 item in cart; on cart page.",
            "Clicking 'Remove' eliminates item row from cart table and removes the shopping cart badge completely.",
            "Item row removed; cart badge disappeared; 0 items remaining.",
            "PASS",
            "None. Cart removal works cleanly.",
            "Medium. Inability to remove unwanted items prevents users from completing desired transactions.",
            "Medium. Cart modification regression.",
            "Automate removing item from inventory page vs removing from cart page.",
            "evidence/saucedemo/tc05b_remove_from_cart.png"
        ],
        [
            "TC05c",
            "Empty Cart Default State",
            "Cart / Zero State",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "Logged in with fresh session.",
            "Cart page displays 0 item rows, no phantom badge, and active 'Continue Shopping' CTA.",
            "Cart rendered 0 items with clean layout.",
            "PASS",
            "None. Empty state handled cleanly.",
            "Low. Minor presentation issue if stale items persist.",
            "Low. Verifies clean initial session state.",
            "Automate 'Continue Shopping' button redirection back to /inventory.html.",
            "evidence/saucedemo/tc05c_empty_cart.png"
        ],
        [
            "TC06",
            "End-to-End Checkout Happy Path",
            "Checkout / Complete Order Funnel",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Smoke",
            "2 items in cart on cart.html.",
            "Complete 3-step checkout: enter customer info -> verify subtotal ($39.98), tax ($3.20), total ($43.18) -> finish order -> confirmation header 'Thank you for your order!'.",
            "Step 1 validated; Step 2 overview confirmed item total $39.98 + tax $3.20 = total $43.18; Step 3 confirmed with 'Thank you for your order!'.",
            "PASS",
            "None. Full checkout pipeline verified.",
            "High. Checkout is the primary revenue-generating flow of the entire application.",
            "Critical. Any defect here directly blocks order placement and revenue.",
            "Automate calculation accuracy across tax bracket variations and discount code inputs.",
            "evidence/saucedemo/tc06_checkout_step2_overview.png\nevidence/saucedemo/tc06_checkout_complete.png"
        ],
        [
            "TC06b",
            "Checkout Step 1 Required Field Validation",
            "Checkout / Form Validation",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Functional",
            "On checkout-step-one.html.",
            "Clicking 'Continue' with empty fields displays 'Error: First Name is required' and blocks step 2 navigation.",
            "Error banner displayed 'Error: First Name is required'; remained on step 1.",
            "PASS",
            "None. Mandatory form fields enforced.",
            "Medium. Prevents incomplete or corrupt shipping records.",
            "Medium. Ensures complete shipping address capture.",
            "Automate sequential missing fields: missing Last Name, missing Postal Code.",
            "evidence/saucedemo/tc06b_checkout_validation.png"
        ],
        [
            "TC07",
            "User Logout via Sidebar Navigation",
            "Authentication / Session Termination",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Smoke",
            "Logged in on inventory page.",
            "Clicking burger menu > Logout clears active session and redirects back to base login page with login form visible.",
            "Redirected to https://www.saucedemo.com/; login button visible.",
            "PASS",
            "None. Session terminated cleanly.",
            "High. Inability to log out compromises user security on shared workstations.",
            "High. Core session lifecycle management.",
            "Automate logout from various subpages (cart, checkout, product details).",
            "evidence/saucedemo/tc07_logout.png"
        ],
        [
            "TC07b",
            "Post-Logout Inventory Security Guard",
            "Security / Session Invalidation",
            "Selenium WebDriver (Python)",
            "unittest + Page Object Methods",
            "Regression",
            "Logged out from SauceDemo.",
            "Direct navigation to https://www.saucedemo.com/inventory.html is blocked and redirected to login page.",
            "Direct navigation immediately redirected to https://www.saucedemo.com/.",
            "PASS",
            "None. Route protection active.",
            "High. Security vulnerability if unauthenticated users can access inventory or cart after logout.",
            "High. Route guard authorization regression.",
            "Automate direct URL attempts to cart.html, checkout-step-one.html, and checkout-step-two.html post-logout.",
            "evidence/saucedemo/tc07b_post_logout_session_check.png"
        ]
    ]

    for row_idx, r_data in enumerate(data_s2, 2):
        ws2.row_dimensions[row_idx].height = 65
        for col_idx, val in enumerate(r_data, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_body
            cell.border = border_thin
            if col_idx in [1, 4, 5, 6, 10]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            
            # Status coloring
            if col_idx == 10:
                if val == "PASS":
                    cell.fill = fill_pass
                    cell.font = font_pass
                elif val == "FAIL":
                    cell.fill = fill_fail

    # -------------------------------------------------------------
    # COLUMN WIDTHS & PANE FREEZING
    # -------------------------------------------------------------
    col_widths_s1 = [10, 22, 26, 16, 32, 26, 36, 34, 34, 34, 12, 32, 36, 36]
    for col_idx, width in enumerate(col_widths_s1, 1):
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = width
    ws1.freeze_panes = "A2"
    ws1.auto_filter.ref = f"A1:{get_column_letter(len(headers_s1))}{len(data_s1)+1}"

    col_widths_s2 = [10, 26, 22, 20, 24, 14, 24, 34, 34, 14, 22, 28, 28, 36, 32]
    for col_idx, width in enumerate(col_widths_s2, 1):
        col_letter = get_column_letter(col_idx)
        ws2.column_dimensions[col_letter].width = width
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers_s2))}{len(data_s2)+1}"

    output_path = "QA_Assessment_Report.xlsx"
    wb.save(output_path)
    print(f"Workbook successfully saved to {output_path}")

if __name__ == "__main__":
    create_qa_workbook()
